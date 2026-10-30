#!/usr/bin/env python3
"""
extraccion_comun.py - Parte compartida por los scripts de extraccion, uno por tipo de documento:

  extraer_anteproyectos.py            -> formato 01  Anteproyecto (presentacion del anteproyecto)
  extraer_evaluacion_anteproyecto.py  -> formato 02  Evaluacion del anteproyecto
  extraer_documento_final.py          -> formato 03  Evaluacion del documento final
  extraer_sustentacion.py             -> formato 18  Acta de socializacion o sustentacion
  extraer_<tipo>.py                   -> un script por cada uno de los demas formatos (04, 05, 07-11, 14-16, 19-22, 24, EC01-EC03, M01, P01, P02,
                                         SI01, SI02, SN01, T01, T02, E01, X01, X02): ver TIPOS mas abajo
  extraer_todos.py                    -> reconoce el tipo de cada archivo de una carpeta y extrae cada tipo por separado

Cada script SOLO acepta archivos de su propio tipo: antes de leer un archivo se revisa que sea ese formato; si es de otro tipo (o no es un
formato FUP, o esta sin diligenciar) se informa y NO se procesa. Asi no se mezclan columnas de formatos distintos en una misma tabla.

Carpetas (se crean solas la primera vez):
  entrada\\<tipo>\\   aqui se copian los Excel diligenciados de ese tipo
  salida\\<tipo>\\    aqui quedan los resultados: extraccion_<tipo>_<fecha>_<hora>.xlsx / .csv y registro_<tipo>_<fecha>_<hora>.txt
                     (cada ejecucion crea archivos nuevos; nunca se sobrescribe una extraccion anterior)
"""
import csv
import datetime
import os
import re
import sys
import unicodedata

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import extraer_formato as ef  # noqa: E402  (lee los formatos por colores y arma la tabla)


MIN_DILIGENCIADO = 0.15   # fraccion minima de campos llenos para considerar que el formato fue diligenciado


def _n(s):
    return "".join(ch for ch in unicodedata.normalize("NFD", str(s).upper()) if unicodedata.category(ch) != "Mn")


# tipo -> nombre legible, hoja del formato y textos que debe (o no debe) tener el encabezado
TIPOS = {
    "anteproyecto": dict(nombre="Anteproyecto (formato 01)", hoja=("Anteproyecto", "Formalizacion"), tiene=("ANTEPROYECTO",), no_tiene=("EVALUACION",)),
    "evaluacion_anteproyecto": dict(nombre="Evaluación del anteproyecto (formato 02)", hoja=("Evaluacion_Anteproyecto",), tiene=("EVALUACION DEL ANTEPROYECTO",), no_tiene=()),
    "documento_final": dict(nombre="Evaluación del documento final (formato 03)", hoja=("Evaluacion_Documento_Final",), tiene=("EVALUACION DEL DOCUMENTO FINAL",), no_tiene=("SOCIALIZACION O SUSTENTACION",)),
    "sustentacion": dict(nombre="Acta de socialización o sustentación (formato 18)", hoja=("Acta_Sustentacion",), tiene=("ACTA DE SOCIALIZACION O SUSTENTACION",), no_tiene=()),
    "evaluacion_sustentacion": dict(nombre="Evaluación de la socialización o sustentación (formato 04)", hoja=("Evaluacion_Documento_Final",), tiene=("SOCIALIZACION O SUSTENTACION",), no_tiene=()),
    # Los demas formatos de las opciones de grado (unificados e individuales) se reconocen por el nombre de su hoja, que es unico en cada formato.
    "aval_anteproyecto": dict(nombre="Aval y aprobación del anteproyecto (formato 05)", hoja=("Aprobacion_Anteproyecto",), tiene=(), no_tiene=()),
    "aval_documento_final": dict(nombre="Aval y aprobación del documento final (formato 07)", hoja=("Aprobacion_Documento_Final",), tiene=(), no_tiene=()),
    "inicio_practica": dict(nombre="Verificación previa al inicio de la práctica y ARL (formato 08)", hoja=("Inicio_Practica",), tiene=(), no_tiene=()),
    "acta_director_evaluadores": dict(nombre="Acta de reunión Director–Evaluadores (formato 09)", hoja=("Acta_Director_Evaluadores",), tiene=(), no_tiene=()),
    "informe_avance_practica": dict(nombre="Informe de avance de la práctica (formato 10)", hoja=("Informe_Avance_Practica",), tiene=(), no_tiene=()),
    "carta_arl_empresa": dict(nombre="Carta de compromiso de ARL de la empresa (formato 11)", hoja=("Carta_Compromiso_ARL",), tiene=(), no_tiene=()),
    "asignacion_director_evaluadores": dict(nombre="Asignación de Director y Evaluadores (formato 14)", hoja=("Asignacion_Dir_Evaluadores",), tiene=(), no_tiene=()),
    "inscripcion": dict(nombre="Inscripción de la opción de grado y verificación de requisitos (formato 15)", hoja=("Inscripcion_Opcion_Grado",), tiene=(), no_tiene=()),
    "registro_asesorias": dict(nombre="Registro de asesorías (formato 16)", hoja=("Registro_Asesorias",), tiene=(), no_tiene=()),
    "prorroga": dict(nombre="Solicitud de prórroga (formato 19)", hoja=("Solicitud_Prorroga",), tiene=(), no_tiene=()),
    "novedad": dict(nombre="Solicitud de novedad del trabajo de grado (formato 20)", hoja=("Solicitud_Novedad",), tiene=(), no_tiene=()),
    "mencion_honor": dict(nombre="Solicitud de mención de honor o trabajo meritorio (formato 21)", hoja=("Mencion_Honor",), tiene=(), no_tiene=()),
    "confidencialidad": dict(nombre="Compromiso de confidencialidad y propiedad intelectual (formato 22)", hoja=("Compromiso_Confidencialidad",), tiene=(), no_tiene=()),
    "informe_director": dict(nombre="Informe semestral del Director (formato 24)", hoja=("Informe_Director",), tiene=(), no_tiene=()),
    "carta_empresa_caso": dict(nombre="Carta de aceptación y de recibido de la empresa (Estudio de Caso) (formato EC01)", hoja=("Carta_Empresa",), tiene=(), no_tiene=()),
    "guia_criterios_caso": dict(nombre="Guía de criterios: Estudio de Caso o Proyecto de Investigación (formato EC02)", hoja=("Guia_Criterios",), tiene=(), no_tiene=()),
    "acta_entrega_resultados": dict(nombre="Acta de entrega de resultados y deslinde de responsabilidad (formato EC03)", hoja=("Acta_Entrega",), tiene=(), no_tiene=()),
    "matriz_revision": dict(nombre="Matriz de revisión bibliográfica (Monografía) (formato M01)", hoja=("Matriz_Revision",), tiene=(), no_tiene=()),
    "acta_compromiso_empresa": dict(nombre="Acta de compromiso con la empresa (Práctica extendida) (formato P01)", hoja=("Acta_Compromiso",), tiene=(), no_tiene=()),
    "evaluacion_tripartita": dict(nombre="Evaluación tripartita de la práctica (formato P02)", hoja=("Evaluacion_Tripartita",), tiene=(), no_tiene=()),
    "control_modulos": dict(nombre="Control de asistencia y cumplimiento de módulos (Seminario) (formato SI01)", hoja=("Control_Modulos",), tiene=(), no_tiene=()),
    "visto_bueno_articulo": dict(nombre="Visto bueno del tutor al artículo recibido (Seminario) (formato SI02)", hoja=("Visto_Bueno_Articulo",), tiene=(), no_tiene=()),
    "aval_internacionalizacion": dict(nombre="Aval de Internacionalización (Seminario Internacional) (formato SN01)", hoja=("Aval_Internacionalizacion",), tiene=(), no_tiene=()),
    "certificacion_semillero": dict(nombre="Certificación de vinculación al semillero (Trayectoria) (formato T01)", hoja=("Certificacion_Semillero",), tiene=(), no_tiene=()),
    "validacion_producto": dict(nombre="Validación del producto avalado (Trayectoria) (formato T02)", hoja=("Validacion_Producto",), tiene=(), no_tiene=()),
    "validacion_mvp": dict(nombre="Ficha de validación del producto mínimo viable (Emprendimiento) (formato E01)", hoja=("Validacion_MVP",), tiene=(), no_tiene=()),
    "solicitud_exencion": dict(nombre="Solicitud de exención por mérito y constancia (formato X01)", hoja=("Exencion_Merito",), tiene=(), no_tiene=()),
    "ficha_resumen_exencion": dict(nombre="Ficha resumen final de la exención (formato X02)", hoja=("Ficha_Resumen",), tiene=(), no_tiene=()),
}


def identificar(ruta):
    """Devuelve la clave de TIPOS a la que pertenece el archivo, o None si no es un formato conocido."""
    wb = openpyxl.load_workbook(ruta)                      # con formulas: los titulos de 01 y 04 son formulas
    try:
        for clave, t in TIPOS.items():
            hoja = next((h for h in t["hoja"] if h in wb.sheetnames), None)
            if not hoja:
                continue
            ws = wb[hoja]
            enc = _n(" ".join(str(ws.cell(r, c).value) for r in range(1, 6) for c in range(1, 12) if ws.cell(r, c).value))
            if all(_n(x) in enc for x in t["tiene"]) and not any(_n(x) in enc for x in t["no_tiene"]):
                return clave
        return None
    finally:
        wb.close()


# ---------------------------------------------------------------- formatos NUEVOS (sin script propio): extraccion adaptable
# Un formato que todavia no esta en TIPOS se reconoce por su ESTRUCTURA (el diseño FUP: celdas azul claro = etiquetas, crema = datos que se diligencian,
# bandas azules = secciones) y se agrupa por el nombre de su hoja. Asi, cualquier formato nuevo hecho con el mismo diseño se extrae sin escribir otro script.
PREFIJO_NUEVO = "nuevo:"
HOJAS_AUX = ("Instrucciones", "Imagenes", "Listas", "Asignacion")


def hoja_formulario(wb):
    return next((n for n in wb.sheetnames if n not in HOJAS_AUX), None)


def identificar_adaptable(ruta):
    """'nuevo:<hoja>' si el archivo tiene el diseño de los formatos FUP (aunque no este registrado en TIPOS); si no, None."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    try:
        hoja = hoja_formulario(wb)
        if not hoja:
            return None
        ws = wb[hoja]
        etiquetas = entradas = bandas = 0
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), max_col=min(ws.max_column, 8)):
            for c in fila:
                f = ef._rgb(c)
                if f == ef.LABEL and c.value not in (None, ""):
                    etiquetas += 1
                elif f == ef.INPUT:
                    entradas += 1
                elif f == ef.BAND and c.value not in (None, ""):
                    bandas += 1
        return PREFIJO_NUEVO + hoja if (etiquetas >= 3 and entradas >= 3 and bandas >= 1) else None
    finally:
        wb.close()


def identificar_todo(ruta):
    """Tipo registrado (clave de TIPOS) o, si el formato es nuevo, 'nuevo:<hoja>'; None si no es un formato FUP."""
    return identificar(ruta) or identificar_adaptable(ruta)


def nombre_tipo(tipo):
    if tipo in TIPOS:
        return TIPOS[tipo]["nombre"]
    hoja = tipo[len(PREFIJO_NUEVO):].split("#")[0]
    return f"Formato nuevo «{hoja}»" + (" (otra estructura)" if "#" in tipo else " (extracción adaptable)")


def carpeta_tipo(tipo):
    """nombre de carpeta (sin caracteres raros) para guardar los resultados del tipo"""
    if tipo.startswith(PREFIJO_NUEVO):
        return "nuevo_" + re.sub(r"[^A-Za-z0-9_]+", "_", _n(tipo[len(PREFIJO_NUEVO):]).lower()).strip("_")
    return tipo


UMBRAL_ESTRUCTURA = 0.5   # parecido minimo (etiquetas en comun) para considerar que dos archivos de un formato nuevo tienen la misma estructura


def etiquetas(ruta):
    """Conjunto de etiquetas (campos y tablas) del formato: sirve para comparar la ESTRUCTURA de dos archivos, sin importar lo que contienen."""
    d = ef.extraer(ruta)
    campos, tablas = [], []
    for sec, cuerpo in d["secciones"].items():
        c, tb = ef._recorrer(cuerpo, sec)
        campos += c
        tablas += tb
    s = {_n(x[2]).strip() for x in campos if x[2]}
    s |= {_n(x[1]).strip() for x in tablas if x[1]}
    return {x for x in s if x}


def _parecido(a, b):
    return len(a & b) / len(a | b) if (a or b) else 1.0


def clasificar(rutas):
    """{ruta: tipo}. tipo = clave de TIPOS (formato registrado), 'nuevo:<hoja>' (formato nuevo), 'nuevo:<hoja>#2'... (misma hoja pero otra
    estructura), None (no es un formato FUP) o '!motivo' (no se pudo abrir). Todos los archivos de un mismo tipo tienen la misma estructura."""
    tipos = {}
    for r in rutas:
        try:
            tipos[r] = identificar_todo(r)
        except Exception as e:  # noqa: BLE001
            tipos[r] = "!" + str(e)
    por_hoja = {}
    for r, tp in tipos.items():
        if isinstance(tp, str) and tp.startswith(PREFIJO_NUEVO):
            por_hoja.setdefault(tp, []).append(r)
    for tp, lista in por_hoja.items():
        grupos = []   # [(clave, etiquetas de referencia)]
        for r in lista:
            try:
                s = etiquetas(r)
            except Exception:  # noqa: BLE001
                continue
            destino = next((k for k, ref in grupos if _parecido(s, ref) >= UMBRAL_ESTRUCTURA), None)
            if destino is None:
                destino = tp if not grupos else f"{tp}#{len(grupos) + 1}"
                grupos.append((destino, s))
            tipos[r] = destino
    return tipos


def _diligenciado(d):
    """(llenos, total) de los campos de entrada del archivo ya extraido."""
    campos, tablas = [], []
    for sec, cuerpo in d["secciones"].items():
        if sec == "RESULTADO CALCULADO POR EL SCRIPT":
            continue
        c, t = ef._recorrer(cuerpo, sec)
        campos += c
        tablas += t
    llenos = sum(1 for *_, v in campos if v not in (None, "")) + sum(1 for *_, r in tablas if any(any(x not in (None, "") for x in f.values()) for f in r))
    return llenos, len(campos) + len(tablas)


def _nuevo(base, ext):
    ruta = base + ext
    n = 2
    while os.path.exists(ruta):
        ruta = f"{base}_{n}{ext}"
        n += 1
    return ruta


def ejecutar(tipo, argv=None):
    """tipo: clave de TIPOS (formato registrado) o 'nuevo:<hoja>' (formato nuevo: extraccion adaptable, ver extraer_adaptable.py)."""
    adaptable = tipo.startswith(PREFIJO_NUEVO)
    base_tipo = tipo.split("#")[0]
    referencia = None
    t = dict(nombre=nombre_tipo(tipo)) if adaptable else TIPOS[tipo]
    carpeta = carpeta_tipo(tipo)
    arg = (argv if argv is not None else sys.argv[1:])
    entrada = os.path.abspath(arg[0]) if arg else os.path.join(AQUI, "entrada", carpeta)
    salida = os.path.join(AQUI, "salida", carpeta)
    os.makedirs(os.path.join(AQUI, "entrada", carpeta), exist_ok=True)
    os.makedirs(salida, exist_ok=True)
    print(f"Extracción de datos · {t['nombre']}")
    if os.path.isdir(entrada):
        archivos = sorted(os.path.join(entrada, f) for f in os.listdir(entrada) if f.lower().endswith(".xlsx") and not f.startswith("~$"))
    elif os.path.isfile(entrada):
        archivos = [entrada]
    else:
        sys.exit(f"No existe: {entrada}")
    if not archivos:
        sys.exit(f"No hay archivos .xlsx en {entrada}. Copie ahí los formatos diligenciados de este tipo y vuelva a ejecutar.")
    ok, rechazados = [], []
    for f in archivos:
        nom = os.path.basename(f)
        try:
            real = identificar_todo(f) if adaptable else identificar(f)
            nuevo = None if (adaptable or real is not None) else identificar_adaptable(f)
            if adaptable and real == base_tipo:
                et = etiquetas(f)
                if referencia is None:
                    referencia = et
                elif _parecido(et, referencia) < UMBRAL_ESTRUCTURA:
                    rechazados.append((nom, "no tiene la misma estructura que los demás archivos de este formato"))
                    continue
        except Exception as e:  # noqa: BLE001
            rechazados.append((nom, f"no se pudo abrir ({e})"))
            continue
        if real != (base_tipo if adaptable else tipo):
            if real is None and nuevo:
                rechazados.append((nom, "es un formato nuevo, sin script propio: use extraer_adaptable.py"))
            else:
                rechazados.append((nom, "no es un formato FUP reconocido" if real is None else ("es de otro tipo: " + nombre_tipo(real))))
            continue
        try:
            d = ef.extraer(f)
        except Exception as e:  # noqa: BLE001
            rechazados.append((nom, f"error al leerlo ({e})"))
            continue
        llenos, total = _diligenciado(d)
        if total and llenos / total < MIN_DILIGENCIADO:      # los formatos en blanco traen unos pocos datos fijos (facultad, programa)
            rechazados.append((nom, f"está sin diligenciar ({llenos} de {total} campos; se exige al menos {MIN_DILIGENCIADO:.0%})"))
            continue
        d["_tipo_extraccion"] = tipo
        ok.append((d, llenos, total))
    ahora = datetime.datetime.now()
    base = os.path.join(salida, f"extraccion_{carpeta}_{ahora:%Y%m%d_%H%M%S}")
    lineas = [f"Extracción · {t['nombre']} · {ahora:%d/%m/%Y %H:%M}", f"Carpeta leída: {entrada}", ""]
    if ok:
        filas = [ef.fila_csv(d) for d, *_ in ok]
        cols = list(dict.fromkeys(k for fila in filas for k in fila))
        cols = ["Archivo", "Tipo de documento", "Fecha de extracción", "% diligenciado"] + [c for c in cols if c != "Archivo"]
        for fila, (d, llenos, total) in zip(filas, ok):
            fila["Tipo de documento"] = t["nombre"]
            fila["Fecha de extracción"] = ahora.strftime("%Y-%m-%d")
            fila["% diligenciado"] = round(llenos / total, 2) if total else None
        ruta_csv = _nuevo(base, ".csv")
        with open(ruta_csv, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=cols, delimiter=";", extrasaction="ignore")
            w.writeheader()
            w.writerows(filas)
        ruta_x = _nuevo(base, ".xlsx")
        ef.a_excel([d for d, *_ in ok], ruta_x)
        print(f"  Procesados: {len(ok)}  ->  {ruta_x}\n                     {ruta_csv}")
        lineas += [f"Procesados ({len(ok)}):"] + [f"  {d['_archivo']}  ({llenos} de {total} campos diligenciados)" for d, llenos, total in ok]
    else:
        print("  No se procesó ningún archivo.")
        lineas.append("No se procesó ningún archivo.")
    if rechazados:
        print(f"  No procesados ({len(rechazados)}):")
        for nom, motivo in rechazados:
            print(f"    - {nom}: {motivo}")
        lineas += ["", f"No procesados ({len(rechazados)}):"] + [f"  {nom}: {motivo}" for nom, motivo in rechazados]
    with open(_nuevo(os.path.join(salida, f"registro_{carpeta}_{ahora:%Y%m%d_%H%M%S}"), ".txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lineas) + "\n")
    return len(ok), rechazados
